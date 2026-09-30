"""
BearerMail Drive, storage quota and share links.

* Drive: each mailbox has its own files and folders. Files live in MongoDB (collections drive_files and
  drive_chunks, 1 MB per chunk), so the usual `mongodump` backup includes them.
* Save to Drive: an email as PDF, Word (.docx), .eml or text, or one of its attachments.
* Storage: mail + Drive count against a per-mailbox quota (DEFAULT_QUOTA_MB, changeable per user;
  0 = unlimited). A full mailbox cannot upload more; incoming mail is still accepted.
* Share links: short codes (/s/<code> on the web app) for a file or a calendar event, with an optional
  password and expiry. Only a bcrypt hash of a link password is stored.

Every endpoint needs the internal API key; the web app decides who may act for which mailbox.
"""

import hashlib
import io
import logging
import os
import re
import secrets
import string
from datetime import datetime, timedelta, timezone
from email import policy
from email.message import EmailMessage
from email.utils import formataddr, format_datetime
from html.parser import HTMLParser
from urllib.parse import quote

import bcrypt
from bson import ObjectId
from fastapi import APIRouter, Body, Depends, HTTPException, Request
from fastapi.responses import Response, StreamingResponse

import bearer_ext as ext
import users_ext

logger = logging.getLogger("bearermail.drive")

CHUNK = 1024 * 1024
DEFAULT_QUOTA_MB = max(0, int(os.getenv("DEFAULT_QUOTA_MB", "5120")))
DRIVE_MAX_FILE_MB = max(1, int(os.getenv("DRIVE_MAX_FILE_MB", "100")))
ATTACH_FROM_DRIVE_MAX_MB = max(1, int(os.getenv("ATTACH_FROM_DRIVE_MAX_MB", "20")))
MAX_FOLDER_DEPTH = 10

_cfg = {"get_db": None, "require_api_key": None}


def configure(get_db, require_api_key):
    _cfg["get_db"] = get_db
    _cfg["require_api_key"] = require_api_key


def _db():
    return _cfg["get_db"]()


def _auth(request: Request):
    _cfg["require_api_key"](request)


router = APIRouter(dependencies=[Depends(_auth)])


def _now():
    return datetime.now(timezone.utc)


def init_indexes():
    db = _db()
    db.drive_files.create_index([("owner", 1), ("folder", 1)])
    db.drive_chunks.create_index([("file_id", 1), ("n", 1)], unique=True)
    db.drive_folders.create_index([("owner", 1), ("path", 1)], unique=True)
    db.shares.create_index("code", unique=True)
    db.shares.create_index("owner")


# ---------------------------------------------------------------------------
# Owners, names and folders
# ---------------------------------------------------------------------------

def owner_of(address: str) -> dict:
    """The mailbox behind an address (a mailbox itself, or the mailbox an alias delivers to)."""
    address = ext._clean_address(address)
    db = _db()
    acc = db.accounts.find_one({"address": address}) if address else None
    if not acc and address:
        alias = db.aliases.find_one({"address": address})
        if alias:
            acc = db.accounts.find_one({"address": alias.get("deliver_to")})
    if not acc:
        raise HTTPException(status_code=404, detail="No mailbox with that address")
    return acc


_BAD_NAME = re.compile(r"[\x00-\x1f\x7f/\\]")


def clean_name(value, fallback: str = "") -> str:
    name = _BAD_NAME.sub("", str(value or "")).strip()[:200]
    if name in ("", ".", ".."):
        if fallback:
            return fallback
        raise HTTPException(status_code=422, detail="Enter a name")
    return name


def clean_folder(value) -> str:
    parts = [p for p in str(value or "/").split("/") if p.strip()]
    if len(parts) > MAX_FOLDER_DEPTH:
        raise HTTPException(status_code=422, detail="Folders can be at most 10 levels deep")
    out = []
    for p in parts:
        p = _BAD_NAME.sub("", p).strip()[:100]
        if p in ("", ".", ".."):
            raise HTTPException(status_code=422, detail="That folder name is not allowed")
        out.append(p)
    return "/" + "/".join(out) if out else "/"


def _unique_name(owner: str, folder: str, name: str, skip_id=None) -> str:
    db = _db()
    q = {"owner": owner, "folder": folder}
    taken = {f["name"].lower() for f in db.drive_files.find(q, {"name": 1}) if f["_id"] != skip_id}
    if name.lower() not in taken:
        return name
    stem, dot, ext_ = name.rpartition(".")
    if not dot or not stem:
        stem, ext_ = name, ""
    for i in range(1, 1000):
        candidate = f"{stem} ({i}){'.' + ext_ if ext_ else ''}"
        if candidate.lower() not in taken:
            return candidate
    raise HTTPException(status_code=409, detail="Too many files with that name")


def _file(owner: str, file_id: str) -> dict:
    try:
        f = _db().drive_files.find_one({"_id": ObjectId(file_id), "owner": owner})
    except Exception:
        f = None
    if not f:
        raise HTTPException(status_code=404, detail="File not found")
    return f


def _public_file(f: dict) -> dict:
    return {"id": str(f["_id"]), "name": f["name"], "folder": f["folder"], "size": f.get("size", 0),
            "content_type": f.get("content_type", "application/octet-stream"), "created_at": ext._iso(f.get("created_at")),
            "updated_at": ext._iso(f.get("updated_at")), "source": f.get("source", {}).get("kind", "upload"),
            "shared": bool(_db().shares.count_documents({"kind": "file", "target_id": str(f["_id"]), "revoked": {"$ne": True}}, limit=1))}


# ---------------------------------------------------------------------------
# Storage and quota
# ---------------------------------------------------------------------------

def _sum(coll, match: dict, field: str) -> int:
    rows = list(coll.aggregate([{"$match": match}, {"$group": {"_id": None, "total": {"$sum": f"${field}"}}}]))
    return int(rows[0]["total"]) if rows else 0


def quota_bytes(acc: dict) -> int:
    mb = acc.get("quota_mb")
    if not isinstance(mb, int) or mb < 0:
        mb = DEFAULT_QUOTA_MB
    return mb * 1024 * 1024


def usage(acc: dict) -> dict:
    db = _db()
    addresses = users_ext._addresses_for(acc["address"])
    mail = _sum(db.messages, {"to_addresses": {"$in": addresses}}, "size")
    drive = _sum(db.drive_files, {"owner": acc["address"]}, "size")
    quota = quota_bytes(acc)
    used = mail + drive
    return {"address": acc["address"], "mail": mail, "drive": drive, "used": used, "quota": quota,
            "percent": round(used * 100 / quota, 1) if quota else 0, "unlimited": quota == 0,
            "full": bool(quota) and used >= quota}


@router.get("/admin/storage/{address}")
def get_storage(address: str):
    return usage(owner_of(address))


@router.post("/admin/storage/{address}/quota")
def set_quota(address: str, body: dict = Body(...)):
    acc = owner_of(address)
    mb = body.get("quota_mb")
    if mb is None:
        _db().accounts.update_one({"_id": acc["_id"]}, {"$unset": {"quota_mb": ""}})
    else:
        if not isinstance(mb, int) or not 0 <= mb <= 10_000_000:
            raise HTTPException(status_code=422, detail="The quota must be a number of MB (0 = unlimited)")
        _db().accounts.update_one({"_id": acc["_id"]}, {"$set": {"quota_mb": mb}})
    return usage(_db().accounts.find_one({"_id": acc["_id"]}))


def _room_for(acc: dict, extra: int):
    u = usage(acc)
    if u["quota"] and u["used"] + extra > u["quota"]:
        raise HTTPException(status_code=507, detail="Your storage is full. Delete files or old mail, or ask your admin for more space.")
    return u


# ---------------------------------------------------------------------------
# Listing, folders
# ---------------------------------------------------------------------------

@router.get("/admin/drive/{address}")
def list_drive(address: str, folder: str = "/"):
    acc = owner_of(address)
    owner = acc["address"]
    folder = clean_folder(folder)
    db = _db()
    prefix = "" if folder == "/" else folder
    sub = set()
    pattern = "^" + re.escape(prefix + "/") + "[^/]+"
    for doc in db.drive_folders.find({"owner": owner, "path": {"$regex": pattern}}, {"path": 1}):
        sub.add(doc["path"][len(prefix) + 1:].split("/")[0])
    for doc in db.drive_files.find({"owner": owner, "folder": {"$regex": pattern}}, {"folder": 1}):
        sub.add(doc["folder"][len(prefix) + 1:].split("/")[0])
    files = [_public_file(f) for f in db.drive_files.find({"owner": owner, "folder": folder}).sort("name", 1)]
    parts = [p for p in folder.split("/") if p]
    crumbs = [{"name": "My Drive", "path": "/"}] + [{"name": p, "path": "/" + "/".join(parts[:i + 1])} for i, p in enumerate(parts)]
    return {"owner": owner, "folder": folder, "breadcrumbs": crumbs,
            "folders": [{"name": n, "path": (prefix + "/" + n)} for n in sorted(sub, key=str.lower)],
            "files": files, "storage": usage(acc), "max_file_mb": DRIVE_MAX_FILE_MB}


@router.get("/admin/drive/{address}/all")
def all_files(address: str, q: str = ""):
    """Every file (for the Compose picker), newest first, optionally filtered by name."""
    owner = owner_of(address)["address"]
    query = {"owner": owner}
    if q:
        query["name"] = {"$regex": re.escape(q[:100]), "$options": "i"}
    return {"files": [_public_file(f) for f in _db().drive_files.find(query).sort("updated_at", -1).limit(200)]}


@router.post("/admin/drive/{address}/folders", status_code=201)
def create_folder(address: str, body: dict = Body(...)):
    owner = owner_of(address)["address"]
    parent = clean_folder(body.get("folder"))
    name = clean_name(body.get("name"))
    path = clean_folder(parent.rstrip("/") + "/" + name)
    _db().drive_folders.update_one({"owner": owner, "path": path}, {"$setOnInsert": {"owner": owner, "path": path, "created_at": _now()}}, upsert=True)
    return {"path": path}


@router.post("/admin/drive/{address}/folders/rename")
def rename_folder(address: str, body: dict = Body(...)):
    owner = owner_of(address)["address"]
    old = clean_folder(body.get("path"))
    if old == "/":
        raise HTTPException(status_code=422, detail="My Drive cannot be renamed")
    new = clean_folder(old.rsplit("/", 1)[0] + "/" + clean_name(body.get("name")))
    db = _db()
    rx = {"$regex": "^" + re.escape(old) + "(/|$)"}
    for doc in list(db.drive_folders.find({"owner": owner, "path": rx})):
        db.drive_folders.update_one({"_id": doc["_id"]}, {"$set": {"path": new + doc["path"][len(old):]}})
    for doc in list(db.drive_files.find({"owner": owner, "folder": rx}, {"folder": 1})):
        db.drive_files.update_one({"_id": doc["_id"]}, {"$set": {"folder": new + doc["folder"][len(old):]}})
    return {"path": new}


@router.delete("/admin/drive/{address}/folders")
def delete_folder(address: str, path: str, recursive: int = 0):
    owner = owner_of(address)["address"]
    path = clean_folder(path)
    if path == "/":
        raise HTTPException(status_code=422, detail="My Drive cannot be deleted")
    db = _db()
    rx = {"$regex": "^" + re.escape(path) + "(/|$)"}
    files = list(db.drive_files.find({"owner": owner, "folder": rx}, {"_id": 1}))
    if files and not recursive:
        raise HTTPException(status_code=409, detail="The folder is not empty")
    for f in files:
        _delete_file_data(f["_id"])
    db.drive_folders.delete_many({"owner": owner, "path": rx})
    return {"deleted_files": len(files)}


# ---------------------------------------------------------------------------
# Files
# ---------------------------------------------------------------------------

def _delete_file_data(file_id):
    db = _db()
    db.drive_chunks.delete_many({"file_id": file_id})
    db.drive_files.delete_one({"_id": file_id})
    db.shares.update_many({"kind": "file", "target_id": str(file_id)}, {"$set": {"revoked": True}})


def store_bytes(acc: dict, folder: str, name: str, data: bytes, content_type: str, source: dict) -> dict:
    """Save a whole file at once (email exports, attachments)."""
    if len(data) > DRIVE_MAX_FILE_MB * 1024 * 1024:
        raise HTTPException(status_code=413, detail=f"Files can be at most {DRIVE_MAX_FILE_MB} MB")
    _room_for(acc, len(data))
    db = _db()
    folder = clean_folder(folder)
    now = _now()
    doc = {"owner": acc["address"], "folder": folder, "name": _unique_name(acc["address"], folder, clean_name(name, "file")),
           "size": len(data), "content_type": (content_type or "application/octet-stream")[:100],
           "sha256": hashlib.sha256(data).hexdigest(), "created_at": now, "updated_at": now, "source": source, "complete": True}
    fid = db.drive_files.insert_one(doc).inserted_id
    for n, i in enumerate(range(0, len(data), CHUNK)):
        db.drive_chunks.insert_one({"file_id": fid, "n": n, "data": data[i:i + CHUNK]})
    doc["_id"] = fid
    return _public_file(doc)


@router.post("/admin/drive/{address}/files", status_code=201)
async def upload(address: str, request: Request, name: str = "", folder: str = "/", content_type: str = "", size: int = 0):
    """Upload a file as the raw request body (streamed, stored in 1 MB chunks)."""
    acc = owner_of(address)
    name = clean_name(name, "file")
    folder = clean_folder(folder)
    limit = DRIVE_MAX_FILE_MB * 1024 * 1024
    declared = max(int(request.headers.get("content-length") or 0), max(0, size))
    if declared > limit:
        raise HTTPException(status_code=413, detail=f"Files can be at most {DRIVE_MAX_FILE_MB} MB")
    u = _room_for(acc, declared)
    db = _db()
    now = _now()
    doc = {"owner": acc["address"], "folder": folder, "name": _unique_name(acc["address"], folder, name), "size": 0,
           "content_type": (content_type or "application/octet-stream")[:100], "created_at": now, "updated_at": now,
           "source": {"kind": "upload"}, "complete": False}
    fid = db.drive_files.insert_one(doc).inserted_id
    digest = hashlib.sha256()
    buf = bytearray()
    size = 0
    n = 0
    try:
        async for piece in request.stream():
            size += len(piece)
            if size > limit:
                raise HTTPException(status_code=413, detail=f"Files can be at most {DRIVE_MAX_FILE_MB} MB")
            if u["quota"] and u["used"] + size > u["quota"]:
                raise HTTPException(status_code=507, detail="Your storage is full")
            digest.update(piece)
            buf.extend(piece)
            while len(buf) >= CHUNK:
                db.drive_chunks.insert_one({"file_id": fid, "n": n, "data": bytes(buf[:CHUNK])})
                del buf[:CHUNK]
                n += 1
        if buf:
            db.drive_chunks.insert_one({"file_id": fid, "n": n, "data": bytes(buf)})
    except BaseException:
        _delete_file_data(fid)
        raise
    db.drive_files.update_one({"_id": fid}, {"$set": {"size": size, "sha256": digest.hexdigest(), "complete": True}})
    return _public_file(db.drive_files.find_one({"_id": fid}))


def file_bytes(owner: str, file_id: str) -> tuple[dict, bytes]:
    f = _file(owner, file_id)
    data = b"".join(c["data"] for c in _db().drive_chunks.find({"file_id": f["_id"]}).sort("n", 1))
    return f, data


def _disposition(name: str, inline: bool = False) -> str:
    ascii_name = name.encode("ascii", "ignore").decode().replace('"', "") or "file"
    return f'{"inline" if inline else "attachment"}; filename="{ascii_name}"; filename*=UTF-8\'\'{quote(name)}'


def stream_file(f: dict, inline: bool = False) -> StreamingResponse:
    def chunks():
        for c in _db().drive_chunks.find({"file_id": f["_id"]}).sort("n", 1):
            yield c["data"]
    return StreamingResponse(chunks(), media_type=f.get("content_type") or "application/octet-stream",
                             headers={"Content-Disposition": _disposition(f["name"], inline), "Content-Length": str(f.get("size", 0))})


@router.get("/admin/drive/{address}/files/{file_id}")
def file_meta(address: str, file_id: str):
    return _public_file(_file(owner_of(address)["address"], file_id))


@router.get("/admin/drive/{address}/files/{file_id}/content")
def file_content(address: str, file_id: str, inline: int = 0):
    return stream_file(_file(owner_of(address)["address"], file_id), bool(inline))


@router.patch("/admin/drive/{address}/files/{file_id}")
def update_file(address: str, file_id: str, body: dict = Body(...)):
    owner = owner_of(address)["address"]
    f = _file(owner, file_id)
    folder = clean_folder(body["folder"]) if "folder" in body else f["folder"]
    name = clean_name(body["name"]) if "name" in body else f["name"]
    name = _unique_name(owner, folder, name, skip_id=f["_id"])
    _db().drive_files.update_one({"_id": f["_id"]}, {"$set": {"folder": folder, "name": name, "updated_at": _now()}})
    return _public_file(_db().drive_files.find_one({"_id": f["_id"]}))


@router.delete("/admin/drive/{address}/files/{file_id}")
def delete_file(address: str, file_id: str):
    f = _file(owner_of(address)["address"], file_id)
    _delete_file_data(f["_id"])
    return {"ok": True}


# ---------------------------------------------------------------------------
# Save an email (or an attachment) to Drive
# ---------------------------------------------------------------------------

class _TextOf(HTMLParser):
    BLOCK = {"p", "div", "br", "tr", "li", "h1", "h2", "h3", "h4", "h5", "h6", "table", "blockquote", "hr"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("style", "script", "head", "title"):
            self.skip += 1
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in ("style", "script", "head", "title"):
            self.skip = max(0, self.skip - 1)
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data)


def html_to_text(html: str) -> str:
    p = _TextOf()
    try:
        p.feed(html or "")
        p.close()
    except Exception:
        return re.sub(r"<[^>]+>", " ", html or "")
    text = "".join(p.out)
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    return re.sub(r"\n\s*\n\s*\n+", "\n\n", text).strip()


def _message_for(acc: dict, message_id: str, sent: bool = False) -> dict:
    """An inbox message (or a Sent record) that belongs to this mailbox."""
    db = _db()
    try:
        oid = ObjectId(message_id)
    except Exception:
        raise HTTPException(status_code=404, detail="Message not found")
    addresses = users_ext._addresses_for(acc["address"])
    if sent:
        doc = db.sent_messages.find_one({"_id": oid, "$or": [{"owner": acc["address"]}, {"from_address": {"$in": addresses}}]})
        if doc:
            doc = {"subject": doc.get("subject", ""), "from": {"address": doc.get("from_address", ""), "name": ""},
                   "to": [{"address": a, "name": ""} for a in doc.get("to", [])], "text": doc.get("text", ""),
                   "html": doc.get("html", ""), "created_at": doc.get("created_at"), "attachments": []}
    else:
        doc = db.messages.find_one({"_id": oid, "to_addresses": {"$in": addresses}})
    if not doc:
        raise HTTPException(status_code=404, detail="Message not found")
    return doc


def _addr(entry) -> str:
    if isinstance(entry, dict):
        return formataddr((entry.get("name") or "", entry.get("address") or "")) if entry.get("name") else (entry.get("address") or "")
    return str(entry or "")


def _when(dt) -> datetime:
    if isinstance(dt, datetime):
        return dt.replace(tzinfo=dt.tzinfo or timezone.utc)
    return _now()


def message_summary(doc: dict) -> dict:
    body = doc.get("text") or html_to_text(doc.get("html", ""))
    return {"subject": doc.get("subject") or "(no subject)", "from": _addr(doc.get("from")),
            "to": ", ".join(_addr(t) for t in (doc.get("to") or [])), "date": _when(doc.get("created_at")),
            "body": body or "", "attachments": [a.get("filename", "attachment") for a in doc.get("attachments") or []]}


_FONT_DIRS = ["/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu", os.path.join(os.path.dirname(__file__), "fonts")]


def _pdf_font(pdf) -> str:
    for d in _FONT_DIRS:
        regular, bold = os.path.join(d, "DejaVuSans.ttf"), os.path.join(d, "DejaVuSans-Bold.ttf")
        if os.path.exists(regular) and os.path.exists(bold):
            pdf.add_font("DejaVu", "", regular)
            pdf.add_font("DejaVu", "B", bold)
            return "DejaVu"
    return ""


def to_pdf(doc: dict) -> bytes:
    from fpdf import FPDF
    s = message_summary(doc)
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    font = _pdf_font(pdf)
    if font:
        clean = str
    else:  # the built-in fonts only know Latin-1
        font = "Helvetica"

        def clean(v):
            return str(v).encode("latin-1", "replace").decode("latin-1")
    pdf.set_title(clean(s["subject"]))
    pdf.set_font(font, "B", 15)
    pdf.multi_cell(0, 8, clean(s["subject"]), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font(font, "", 10)
    for label, value in (("From", s["from"]), ("To", s["to"]), ("Date", s["date"].strftime("%Y-%m-%d %H:%M UTC")),
                         ("Attachments", ", ".join(s["attachments"]))):
        if value:
            pdf.set_font(font, "B", 10)
            pdf.cell(30, 6, clean(label + ":"))
            pdf.set_font(font, "", 10)
            pdf.multi_cell(0, 6, clean(value), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_draw_color(180, 180, 180)
    pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
    pdf.ln(4)
    pdf.set_font(font, "", 11)
    pdf.multi_cell(0, 6, clean(s["body"] or "(empty message)"), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def to_docx(doc: dict) -> bytes:
    import docx
    s = message_summary(doc)
    d = docx.Document()
    d.core_properties.title = s["subject"][:200]
    d.add_heading(s["subject"], level=1)
    for label, value in (("From", s["from"]), ("To", s["to"]), ("Date", s["date"].strftime("%Y-%m-%d %H:%M UTC")),
                         ("Attachments", ", ".join(s["attachments"]))):
        if value:
            p = d.add_paragraph()
            p.add_run(label + ": ").bold = True
            p.add_run(value)
    for block in (s["body"] or "(empty message)").split("\n\n"):
        d.add_paragraph(block)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def to_eml(doc: dict) -> bytes:
    msg = EmailMessage(policy=policy.SMTP)
    msg["From"] = _addr(doc.get("from")) or "unknown@invalid"
    to = ", ".join(_addr(t) for t in (doc.get("to") or []))
    if to:
        msg["To"] = to
    msg["Subject"] = doc.get("subject") or ""
    msg["Date"] = format_datetime(_when(doc.get("created_at")))
    text, html = doc.get("text") or "", doc.get("html") or ""
    msg.set_content(text or html_to_text(html) or " ")
    if html:
        msg.add_alternative(html, subtype="html")
    for att in doc.get("attachments") or []:
        content = att.get("content")
        if isinstance(content, (bytes, bytearray)):
            maintype, _, subtype = (att.get("content_type") or "application/octet-stream").partition("/")
            msg.add_attachment(bytes(content), maintype=maintype or "application", subtype=subtype or "octet-stream",
                               filename=att.get("filename") or "attachment")
    return msg.as_bytes()


def to_txt(doc: dict) -> bytes:
    s = message_summary(doc)
    head = [f"Subject: {s['subject']}", f"From: {s['from']}", f"To: {s['to']}", f"Date: {s['date']:%Y-%m-%d %H:%M UTC}"]
    if s["attachments"]:
        head.append("Attachments: " + ", ".join(s["attachments"]))
    return ("\n".join(head) + "\n\n" + s["body"] + "\n").encode("utf-8")


FORMATS = {
    "pdf": (to_pdf, "application/pdf", "pdf"),
    "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "eml": (to_eml, "message/rfc822", "eml"),
    "txt": (to_txt, "text/plain; charset=utf-8", "txt"),
}


def _safe_filename(subject: str) -> str:
    base = re.sub(r"[\x00-\x1f\x7f/\\:*?\"<>|]+", " ", subject or "Email").strip()[:80]
    return base or "Email"


@router.post("/admin/drive/{address}/from-message", status_code=201)
def save_message(address: str, body: dict = Body(...)):
    acc = owner_of(address)
    fmt = str(body.get("format") or "pdf").lower()
    if fmt not in FORMATS:
        raise HTTPException(status_code=422, detail="Choose pdf, docx, eml or txt")
    doc = _message_for(acc, str(body.get("message_id") or ""), sent=bool(body.get("sent")))
    render, ctype, extension = FORMATS[fmt]
    data = render(doc)
    return store_bytes(acc, body.get("folder") or "/Saved emails", f"{_safe_filename(doc.get('subject'))}.{extension}",
                       data, ctype, {"kind": "message", "message_id": str(body.get("message_id")), "format": fmt})


@router.post("/admin/drive/{address}/from-attachment", status_code=201)
def save_attachment(address: str, body: dict = Body(...)):
    acc = owner_of(address)
    doc = _message_for(acc, str(body.get("message_id") or ""))
    wanted = str(body.get("attachment_id") or "")
    for i, att in enumerate(doc.get("attachments") or []):
        if wanted in (str(att.get("id")), str(i)):
            content = att.get("content")
            if not isinstance(content, (bytes, bytearray)):
                raise HTTPException(status_code=404, detail="The attachment content is not stored")
            return store_bytes(acc, body.get("folder") or "/Attachments", att.get("filename") or f"attachment-{i}", bytes(content),
                               att.get("content_type") or "application/octet-stream",
                               {"kind": "attachment", "message_id": str(body.get("message_id"))})
    raise HTTPException(status_code=404, detail="Attachment not found")


def attachments_for_send(owner_address: str, file_ids: list) -> list:
    """Drive files to attach to an outgoing message (they must belong to the sender's mailbox)."""
    out, total = [], 0
    for fid in file_ids or []:
        f, data = file_bytes(owner_address, str(fid))
        total += len(data)
        if total > ATTACH_FROM_DRIVE_MAX_MB * 1024 * 1024:
            raise HTTPException(status_code=413, detail=f"Attachments from Drive can be at most {ATTACH_FROM_DRIVE_MAX_MB} MB in total. Share a link instead.")
        out.append({"filename": f["name"], "content_type": (f.get("content_type") or "application/octet-stream").split(";")[0], "raw": data})
    return out


# ---------------------------------------------------------------------------
# Share links
# ---------------------------------------------------------------------------

_CODE_ALPHABET = string.ascii_letters + string.digits


def _new_code() -> str:
    for _ in range(20):
        code = "".join(secrets.choice(_CODE_ALPHABET) for _ in range(8))
        if not _db().shares.find_one({"code": code}):
            return code
    raise HTTPException(status_code=500, detail="Could not create a link, try again")


def _target_name(kind: str, owner: str, target_id: str) -> str:
    db = _db()
    try:
        oid = ObjectId(target_id)
    except Exception:
        return ""
    if kind == "file":
        f = db.drive_files.find_one({"_id": oid, "owner": owner})
        return f["name"] if f else ""
    e = db.events.find_one({"_id": oid, "owner": owner})
    return e.get("title", "") if e else ""


def _public_share(s: dict) -> dict:
    now = _now()
    expires = s.get("expires_at")
    expired = bool(expires and ext._iso(expires) and expires.replace(tzinfo=expires.tzinfo or timezone.utc) < now)
    return {"code": s["code"], "kind": s["kind"], "target_id": s["target_id"], "name": _target_name(s["kind"], s["owner"], s["target_id"]),
            "has_password": bool(s.get("password_hash")), "expires_at": ext._iso(expires), "expired": expired,
            "downloads": s.get("downloads", 0), "views": s.get("views", 0), "created_at": ext._iso(s.get("created_at")),
            "domain": s.get("domain", ""), "web_host": web_host_of(s.get("domain", "")),
            "revoked": bool(s.get("revoked"))}


def _share_domain(acc: dict, wanted) -> str:
    """The domain whose web address a link uses: one the person may use (their mailbox's domain by default)."""
    allowed = users_ext.allowed_domains(acc)
    wanted = ext._clean_address(wanted)
    if wanted:
        if wanted not in allowed:
            raise HTTPException(status_code=403, detail="Your admin has not given you that domain")
        return wanted
    own = acc["address"].split("@", 1)[1]
    return own if own in allowed else (allowed[0] if allowed else own)


def web_host_of(domain: str) -> str | None:
    d = _db().domains.find_one({"domain": domain, "is_active": True}, {"web_host": 1}) if domain else None
    return (d or {}).get("web_host") or None


@router.get("/admin/drive/{address}/share-hosts")
def share_hosts(address: str):
    """Addresses a link can use: the allowed domains that have a web address set (Setup > Domains)."""
    acc = owner_of(address)
    default = _share_domain(acc, None)
    return {"default": default, "hosts": [{"domain": d, "web_host": web_host_of(d)} for d in users_ext.allowed_domains(acc) if web_host_of(d)]}


@router.post("/admin/drive/{address}/shares", status_code=201)
def create_share(address: str, body: dict = Body(...)):
    acc = owner_of(address)
    if not users_ext.permissions_of(acc).get("share_links", True):
        raise HTTPException(status_code=403, detail="Your admin has not allowed share links")
    kind = body.get("kind")
    if kind not in ("file", "event"):
        raise HTTPException(status_code=422, detail="kind must be file or event")
    target_id = str(body.get("target_id") or "")
    if not _target_name(kind, acc["address"], target_id):
        raise HTTPException(status_code=404, detail="Not found")
    password = body.get("password") or ""
    if password and (not isinstance(password, str) or len(password) < 4 or len(password) > 200):
        raise HTTPException(status_code=422, detail="A link password needs 4 to 200 characters")
    days = body.get("expires_days")
    expires = None
    if days not in (None, "", 0):
        if not isinstance(days, int) or not 1 <= days <= 3650:
            raise HTTPException(status_code=422, detail="Expiry must be 1 to 3650 days")
        expires = _now() + timedelta(days=days)
    doc = {"code": _new_code(), "owner": acc["address"], "kind": kind, "target_id": target_id,
           "domain": _share_domain(acc, body.get("domain")),
           "password_hash": bcrypt.hashpw(password.encode(), bcrypt.gensalt()).decode() if password else None,
           "expires_at": expires, "created_at": _now(), "downloads": 0, "views": 0, "revoked": False}
    _db().shares.insert_one(doc)
    return _public_share(doc)


@router.get("/admin/drive/{address}/shares")
def list_shares(address: str):
    owner = owner_of(address)["address"]
    return {"shares": [_public_share(s) for s in _db().shares.find({"owner": owner, "revoked": {"$ne": True}}).sort("created_at", -1)]}


@router.delete("/admin/drive/{address}/shares/{code}")
def delete_share(address: str, code: str):
    owner = owner_of(address)["address"]
    res = _db().shares.update_one({"code": code, "owner": owner}, {"$set": {"revoked": True}})
    if not res.matched_count:
        raise HTTPException(status_code=404, detail="Link not found")
    return {"ok": True}


def _live_share(code: str) -> dict:
    if not re.fullmatch(r"[A-Za-z0-9]{6,16}", code or ""):
        raise HTTPException(status_code=404, detail="This link does not exist")
    s = _db().shares.find_one({"code": code, "revoked": {"$ne": True}})
    if not s:
        raise HTTPException(status_code=404, detail="This link does not exist or was turned off")
    info = _public_share(s)
    if info["expired"]:
        raise HTTPException(status_code=410, detail="This link has expired")
    if not info["name"]:
        raise HTTPException(status_code=404, detail="The shared item was deleted")
    return s


@router.get("/admin/shares/{code}")
def resolve_share(code: str):
    """What a public link points to (used by the web app's /s/<code> page)."""
    s = _live_share(code)
    info = _public_share(s)
    info.pop("target_id", None)
    owner_acc = _db().accounts.find_one({"address": s["owner"]}) or {}
    info["shared_by"] = owner_acc.get("display_name") or s["owner"]
    if s["kind"] == "file":
        f = _db().drive_files.find_one({"_id": ObjectId(s["target_id"])})
        info.update(size=f.get("size", 0), content_type=f.get("content_type", ""))
    else:
        import calendar_ext
        info["event"] = calendar_ext.public_event(_db().events.find_one({"_id": ObjectId(s["target_id"])}))
    return info


@router.post("/admin/shares/{code}/check")
def check_share_password(code: str, body: dict = Body(default={})):
    s = _live_share(code)
    if not s.get("password_hash"):
        return {"ok": True}
    pw = (body or {}).get("password")
    if not isinstance(pw, str) or not bcrypt.checkpw(pw.encode()[:1024], s["password_hash"].encode()):
        raise HTTPException(status_code=403, detail="Wrong password")
    return {"ok": True}


@router.post("/admin/shares/{code}/viewed")
def share_viewed(code: str):
    _db().shares.update_one({"code": code}, {"$inc": {"views": 1}})
    return {"ok": True}


@router.get("/admin/shares/{code}/content")
def share_content(code: str):
    """The shared file or event (.ics). The web app checks the link password before calling this."""
    s = _live_share(code)
    _db().shares.update_one({"_id": s["_id"]}, {"$inc": {"downloads": 1}})
    if s["kind"] == "file":
        f = _db().drive_files.find_one({"_id": ObjectId(s["target_id"])})
        return stream_file(f)
    import calendar_ext
    e = _db().events.find_one({"_id": ObjectId(s["target_id"])})
    return Response(calendar_ext.event_ics(e), media_type="text/calendar; charset=utf-8",
                    headers={"Content-Disposition": _disposition(_safe_filename(e.get("title")) + ".ics")})
